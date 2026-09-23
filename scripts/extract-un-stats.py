import csv
import gzip
import json
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl

source = Path.home() / 'Documents/New project 8/data_downloads/un_desa'
destination = Path(sys.argv[1])
destination.mkdir(parents=True, exist_ok=True)
def write(name, data):
    (destination / name).write_text(json.dumps(data, ensure_ascii=False, indent=2) + '\n')

age_path = source / 'wpp2024/raw/WPP2024_PopulationByAge5GroupSex_Medium_20260906.csv.gz'
locations = {'Africa', 'Asia', 'Europe', 'Northern America', 'Latin America and the Caribbean', 'Oceania'}
age = defaultdict(lambda: {'youth': 0, 'working': 0, 'elderly': 0})
with gzip.open(age_path, 'rt', encoding='utf-8-sig') as handle:
    for record in csv.DictReader(handle):
        if record['Time'] != '2025' or record['LocTypeID'] != '2' or record['Location'] not in locations or not record['AgeGrpStart']:
            continue
        start = int(record['AgeGrpStart'])
        age[record['Location']]['youth' if start < 15 else 'working' if start < 65 else 'elderly'] += float(record['PopTotal'])
assert len(age) == 6 and all(sum(values.values()) > 1000 for values in age.values())
write('wpp2024_continent_age_2025.json', dict(age))

demography_path = source / 'wpp2024/raw/WPP2024_Demographic_Indicators_Medium.csv.gz'
history = defaultdict(dict)
with gzip.open(demography_path, 'rt', encoding='utf-8-sig') as handle:
    for record in csv.DictReader(handle):
        if record['ISO3_code'] and record['Time'] in {'1990', '2000', '2010', '2020', '2025'} and record['CNMR']:
            history[record['ISO3_code']][record['Time']] = float(record['CNMR'])
assert len(history) > 190
write('wpp2024_country_migration_rate_history.json', dict(history))

stock_path = source / 'undesa_pd_2020_ims_stock_by_sex_destination_and_origin.xlsx'
book = openpyxl.load_workbook(stock_path, read_only=True, data_only=True)
regions = {}
for record in book['Table 2'].iter_rows(min_row=12, values_only=True):
    if record[3] in {903, 935, 908, 905, 904, 909}:
        regions[str(record[3])] = {'name': record[1].strip(), 'migrants': int(record[4])}
assert len(regions) == 6 and sum(value['migrants'] for value in regions.values()) == 280598105
destinations = {840:'미국',276:'독일',682:'사우디아라비아',36:'오스트레일리아',124:'캐나다'}
origins = defaultdict(list)
for record in book['Table 1'].iter_rows(min_row=12, values_only=True):
    if record[3] in destinations and isinstance(record[6], int) and record[6] < 900 and isinstance(record[13], (int, float)) and record[6] != 630:
        origins[str(record[3])].append({'country': record[5].strip().rstrip('*'), 'people': int(record[13])})
rank = {destinations[int(code)]: sorted(records, key=lambda item: -item['people'])[:5] for code, records in origins.items()}
assert len(rank) == 5
write('un_migrant_stock_2020.json', {'destination_regions': regions, 'origin_rank': rank})
