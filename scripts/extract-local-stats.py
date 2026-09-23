import csv
import json
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl

source = Path.home() / 'Documents/New project 8/data_downloads'
destination = Path(sys.argv[1])
destination.mkdir(parents=True, exist_ok=True)

def write(name, value):
    (destination / name).write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n')

def numeric(value):
    try:
        return float(value)
    except (TypeError, ValueError):
        return None

age_path = source / 'kosis/raw/DT_1B04005N/101_DT_1B04005N_M_2024.csv'
ages = defaultdict(lambda: {'label': '', 'total': 0, 'male': 0, 'female': 0, 'youth': 0, 'working': 0, 'elderly': 0, 'young_male': 0, 'young_female': 0})
with age_path.open(encoding='utf-8-sig') as handle:
    rows = csv.reader(handle)
    next(rows); next(rows); next(rows)
    for fields in rows:
        if len(fields) < 8 or fields[4] != '202412':
            continue
        code = fields[0].strip("'")
        if len(code) != 2:
            continue
        age = fields[2].strip("'")
        if age == '0':
            ages[code]['label'] = fields[1]
            continue
        start = int(age) - 5
        count, male, female = (int(float(fields[i])) for i in (5, 6, 7))
        record = ages[code]
        record['total'] += count
        record['male'] += male
        record['female'] += female
        record['youth' if start < 15 else 'working' if start < 65 else 'elderly'] += count
        if 20 <= start < 40:
            record['young_male'] += male
            record['young_female'] += female
assert len(ages) == 18 and ages['00']['total'] == 51217221
write('kosis_age_sex_202412.json', dict(ages))

industry_path = source / 'kosis/raw/DT_1FS1101/101_DT_1FS1101_Y_2024.csv'
industry = defaultdict(dict)
with industry_path.open(encoding='utf-8-sig') as handle:
    rows = csv.reader(handle)
    next(rows); next(rows); next(rows)
    for fields in rows:
        if len(fields) < 12 or fields[4] != '2024':
            continue
        region = fields[0].strip("'")
        sector = fields[2].strip("'")
        if len(region) != 2 or sector not in ['C'] + [f'C{i}' for i in range(10, 35)]:
            continue
        industry[region][sector] = {'region': fields[1], 'sector': fields[3].strip(), 'establishments': numeric(fields[5]), 'employees': numeric(fields[6]), 'shipments_million_krw': numeric(fields[8])}
assert len(industry) == 18 and 'C' in industry['00']
write('kosis_manufacturing_2024.json', dict(industry))

keei_path = source / 'keei/2025_지역에너지통계연보_2024자료.xlsx'
book = openpyxl.load_workbook(keei_path, read_only=True, data_only=True)
keei = {}
for sheet, key in [('Ⅰ-2', 'production'), ('Ⅰ-3', 'supply')]:
    ws = book[sheet]
    headers = [str(x.value) for x in ws[5][1:7 if key == 'production' else 8]]
    first, last = 254, 270
    data = {}
    for line in ws.iter_rows(min_row=first, max_row=last, values_only=True):
        data[str(line[0]).strip()] = {'total': float(line[1]), **{headers[i - 2]: float(line[i - 1]) for i in range(2, len(headers) + 2)}}
    assert len(data) == 17
    keei[key] = data
    keei[key + '_headers'] = headers
for sheet, key in [('Ⅰ-4', 'regional_supply'), ('Ⅰ-5', 'final_consumption'), ('Ⅴ-1', 'generation')]:
    ws = book[sheet]
    year_row = 39 if key == 'generation' else 34
    names = [x.value for x in ws[4][1:19]]
    amounts = [x.value for x in ws[year_row][1:19]]
    keei[key] = {str(name).strip(): float(value) for name, value in zip(names, amounts) if name and numeric(value) is not None}
write('keei_regional_energy_2024.json', keei)

knrec_path = source / 'knrec/2024_신재생에너지보급통계_통계표1_연도별지역별현황.xlsx'
book = openpyxl.load_workbook(knrec_path, read_only=True, data_only=True)
renewable = {}
for sheet, key, names_row, start_col, total_row in [
    ('2.1.1 에너지생산량(지역별)_2024', 'production_toe', 7, 4, 9),
    ('2.2.1 발전량(지역별)_2024', 'generation_mwh', 5, 5, 7),
]:
    ws = book[sheet]
    names = [ws.cell(names_row, i).value for i in range(start_col, start_col + 18)]
    amounts = [ws.cell(total_row, i).value for i in range(start_col, start_col + 18)]
    renewable[key] = {str(name).strip(): float(value) for name, value in zip(names, amounts)}
write('knrec_regional_renewable_2024.json', renewable)

ei_path = source / 'energy_institute/EI_Statistical_Review_2025_ALL_data.xlsx'
book = openpyxl.load_workbook(ei_path, read_only=True, data_only=True)
ei = {}
for sheet, key, start in [('TES by fuel', 'supply_ej', 8), ('Elec generation by fuel', 'generation_twh', 9)]:
    ws = book[sheet]
    headers = [str(x.value).replace('Renew- ables', 'Renewables') for x in ws[3][start:]]
    records = {}
    for row in ws.iter_rows(min_row=5, values_only=True):
        if not isinstance(row[0], str):
            continue
        values = row[start: start + len(headers)]
        if all(numeric(x) is not None for x in values):
            records[row[0]] = dict(zip(headers, [float(x) for x in values]))
    ei[key] = records
    ei[key + '_headers'] = headers
ws = book['Renewables Generation by Source']
headers = [str(cell.value) for cell in ws[3][6:11]]
ei['renewable_generation_twh'] = {row[0]: dict(zip(headers, [float(x) for x in row[6:11]])) for row in ws.iter_rows(min_row=5, values_only=True)
    if isinstance(row[0], str) and all(numeric(x) is not None for x in row[6:11])}
for sheet, key in [
    ('Oil Production - tonnes', 'oil_production_mt'), ('Oil Consumption - Tonnes', 'oil_consumption_mt'),
    ('Gas Production - Bcm', 'gas_production_bcm'), ('Gas Consumption - Bcm', 'gas_consumption_bcm'),
    ('Coal Production - EJ', 'coal_production_ej'), ('Coal Consumption - EJ', 'coal_consumption_ej'),
]:
    ws = book[sheet]
    head = next(ws.iter_rows(min_row=3, max_row=3, values_only=True))
    year_index = next(i for i, value in enumerate(head) if str(value).strip() == '2024')
    ei[key] = {row[0]: float(row[year_index]) for row in ws.iter_rows(min_row=5, values_only=True)
        if isinstance(row[0], str) and numeric(row[year_index]) is not None}
write('ei_energy_2024.json', ei)

land_path = source / 'kosis/korea_land_farm_2024/raw/kosis_bubble_cultivated_area_536_info_20260822.json'
land_source = json.loads(land_path.read_text())
land = {r['itmNm']: int(r['rn1'].replace(',', '')) for r in land_source['mobInfoNewSeries'] if r['prdDe1'] == '2025'}
assert len(land) == 17
write('kosis_cultivated_area_2025.json', land)

employment_path = source / 'kosis/raw/DT_1DA7E33S_NEW/행정구역_시도__산업별_취업자_2025.csv.csv'
employment = defaultdict(dict)
labels = {'계': 'total', 'A 농업 임업 및 어업(01~03)': 'agriculture', '* 광공업(BC)': 'mining_manufacturing', '* 사회간접자본 및 기타서비스업(D~U)': 'services'}
with employment_path.open(encoding='utf-8-sig') as handle:
    for record in csv.DictReader(handle):
        field = labels.get(record['산업별(1)'])
        if not field:
            continue
        primary, secondary = record['시도별(1)'], record['시도별(2)']
        region = secondary if primary == '전남광주통합특별시' and secondary != '소계' else primary if secondary == '소계' and primary != '전남광주통합특별시' else None
        if region and region != '계':
            employment[region][field] = numeric(record['2025'])
assert len(employment) == 17 and all(len(value) == 4 for value in employment.values())
write('kosis_employment_2025.json', dict(employment))

crop_path = source / 'kosis/browser_extract/crop_area_2025_20260920/source_cells.json'
crop_cells = json.loads(crop_path.read_text())
crop_data = {}
for key in ('rice', 'vegetables', 'facility', 'fruit'):
    source_rows = crop_cells[key]['rows']
    crop_data[key] = {fields[0]: sum(int(value.replace(',', '')) for value in fields[1:]) for fields in source_rows}
utilization_path = source / 'kosis/browser_extract/DT_1ET0040/DT_1ET0040_total_land_utilization_rate_2025_20260904.json'
utilization = json.loads(utilization_path.read_text())
crop_data['total_cultivated_area'] = {record['region']: record['cultivatedAreaHa'] for record in utilization['rows']}
assert all(len(crop_data[key]) == 18 for key in crop_data)
write('kosis_crop_area_2025.json', crop_data)

land_book = openpyxl.load_workbook(source / 'kosis/korea_land_farm_2024/raw/molit_cadastral_yearbook_2024_12/99_summary_2024_12.xlsx', read_only=True, data_only=True)
land_sheet = land_book['(11)전국시군구별면적순위현황']
land_area = defaultdict(float)
for fields in land_sheet.iter_rows(min_row=6, values_only=True):
    if isinstance(fields[0], int) and isinstance(fields[1], str) and numeric(fields[3]) is not None:
        land_area[fields[1].split()[0]] += float(fields[3]) / 1000000
assert len(land_area) == 17 and 100000 < sum(land_area.values()) < 101000
write('molit_province_area_2024.json', dict(land_area))
